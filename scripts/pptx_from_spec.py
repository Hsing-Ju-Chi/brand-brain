#!/usr/bin/env python3
"""spec.json → PPTX：每頁底稿圖鋪滿＋原生卡片框＋icon 貼圖＋原生文字框（中英字型、字距、行距、透明度都照 HTML 量測值）。
用法：python pptx_from_spec.py spec.json out.pptx  [SLIDE_ONLY=n 只做第 n 頁，QA 用]
"""
import sys, os, json, re
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

PX = 9525  # 1px @96dpi = 9525 EMU；投影片 1280×720px = 13.333×7.5in

def rgba(s):
    m = re.match(r"rgba?\(([\d.]+),\s*([\d.]+),\s*([\d.]+)(?:,\s*([\d.]+))?\)", s or "")
    if not m: return None, 1.0
    r, g, b = (int(float(m.group(i))) for i in (1, 2, 3))
    a = float(m.group(4)) if m.group(4) is not None else 1.0
    return RGBColor(r, g, b), a

def set_alpha(color_el_parent, alpha):
    """在 <a:srgbClr> 下加 <a:alpha val=...>（python-pptx 沒有 API）"""
    srgb = color_el_parent.find(qn("a:srgbClr"))
    if srgb is not None and alpha < 1.0:
        al = etree.SubElement(srgb, qn("a:alpha")); al.set("val", str(int(alpha * 100000)))

def add_text(slide, it):
    tb = slide.shapes.add_textbox(Emu(int(it["x"]*PX)), Emu(int(it["y"]*PX)), Emu(int((it["w"]+6)*PX)), Emu(int((it["h"]+4)*PX)))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    color, _ = rgba(it["color"]); alpha = it.get("opacity", 1.0)
    for i, line in enumerate(it["lines"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = Pt(it["size"] * (it.get("lh") or 1.2) * 0.75)  # 固定行距（pt），跟 CSS line-height 的行盒一樣高；比例行距在 PowerPoint 會偏高
        p.alignment = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(it.get("align"), PP_ALIGN.LEFT)
        r = p.add_run(); r.text = line
        f = r.font; f.size = Pt(it["size"] * 0.75); f.bold = it["weight"] >= 600; f.name = it["font"]
        if color is not None: f.color.rgb = color
        rPr = r._r.get_or_add_rPr()
        # 中文字型：latin 之外要另設 ea，否則 PowerPoint 會退成宋體
        ea = etree.SubElement(rPr, qn("a:ea")); ea.set("typeface", it["font"])
        if it.get("ls"): rPr.set("spc", str(int(it["ls"] * 0.75 * 100)))
        if alpha < 1.0:
            solid = rPr.find(qn("a:solidFill"))
            if solid is not None: set_alpha(solid, alpha)
    return tb

def add_box(slide, it):
    radius = it.get("radius", 0) or 0
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius > 0 else MSO_SHAPE.RECTANGLE,
                                   Emu(int(it["x"]*PX)), Emu(int(it["y"]*PX)), Emu(int(it["w"]*PX)), Emu(int(it["h"]*PX)))
    if radius > 0:
        shape.adjustments[0] = min(0.5, radius / min(it["w"], it["h"]))
    shape.shadow.inherit = False
    bg, a = rgba(it.get("bg"))
    if bg is None or a == 0:
        shape.fill.background()
    else:
        shape.fill.solid(); shape.fill.fore_color.rgb = bg
        set_alpha(shape.fill._xPr.find(qn("a:solidFill")), a)
    # 邊框：有四邊框用四邊；只有左邊線（紙本編輯）畫一條線
    bc, ba = rgba(it.get("border")); bw = it.get("bw", 0) or 0
    if bc is not None and bw > 0 and ba > 0:
        shape.line.color.rgb = bc; shape.line.width = Emu(int(bw * PX))
        set_alpha(shape.line._ln.find(qn("a:solidFill")), ba)
    else:
        shape.line.fill.background()
        blc, bla = rgba(it.get("bl")); blw = it.get("blw", 0) or 0
        if blc is not None and blw > 0 and bla > 0:
            ln = slide.shapes.add_connector(1, Emu(int(it["x"]*PX)), Emu(int(it["y"]*PX)), Emu(int(it["x"]*PX)), Emu(int((it["y"]+it["h"])*PX)))
            ln.line.color.rgb = blc; ln.line.width = Emu(int(blw * PX))
            set_alpha(ln.line._ln.find(qn("a:solidFill")), bla)
    # 文字框放在 shape 之上：add_shape 會自帶空 text_frame，清掉避免佔位
    shape.text_frame.text = ""
    return shape

def main(spec_path, out_path):
    spec = json.load(open(spec_path))
    only = os.environ.get("SLIDE_ONLY")
    prs = Presentation(); prs.slide_width = Emu(1280 * PX); prs.slide_height = Emu(720 * PX)
    for i, s in enumerate(spec["slides"]):
        if only and int(only) != i + 1: continue
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(s["bg"], 0, 0, prs.slide_width, prs.slide_height)
        items = s["items"]
        if spec.get("native"):
            for it in [x for x in items if x["kind"] == "box"]: add_box(slide, it)
            for it in [x for x in items if x["kind"] == "icon"]:
                slide.shapes.add_picture(it["png"], Emu(int(it["x"]*PX)), Emu(int(it["y"]*PX)), Emu(int(it["w"]*PX)), Emu(int(it["h"]*PX)))
        for it in [x for x in items if x["kind"] == "txt"]: add_text(slide, it)
    prs.save(out_path); print("saved", out_path, "slides:", len(prs.slides))

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
